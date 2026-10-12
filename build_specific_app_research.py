from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


OUT = Path(r"C:\Users\WIN-10\Documents\Abqoor\Research_Data_Analysis_Dirasti.docx")
FONT = "Arial"
BLACK = "000000"
NAVY = "17365D"
PALE_BLUE = "EAF2F8"
LIGHT_GRAY = "D9D9D9"


def remove_paragraph_borders(ppr):
    border = ppr.find(qn("w:pBdr"))
    if border is not None:
        ppr.remove(border)


def set_rtl_paragraph(paragraph):
    ppr = paragraph._p.get_or_add_pPr()
    bidi = ppr.find(qn("w:bidi"))
    if bidi is None:
        bidi = OxmlElement("w:bidi")
        ppr.append(bidi)
    bidi.set(qn("w:val"), "1")


def set_run(run, *, size=16, bold=False, color=BLACK):
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)
    rpr = run._element.get_or_add_rPr()
    fonts = rpr.get_or_add_rFonts()
    for key in ("ascii", "hAnsi", "eastAsia", "cs"):
        fonts.set(qn(f"w:{key}"), FONT)
    rtl = rpr.find(qn("w:rtl"))
    if rtl is None:
        rtl = OxmlElement("w:rtl")
        rpr.append(rtl)
    rtl.set(qn("w:val"), "1")
    lang = rpr.find(qn("w:lang"))
    if lang is None:
        lang = OxmlElement("w:lang")
        rpr.append(lang)
    lang.set(qn("w:bidi"), "ar-SA")


def paragraph(doc, text, *, size=16, bold=False, align=WD_ALIGN_PARAGRAPH.LEFT,
              before=0, after=8, keep=False, indent=0):
    p = doc.add_paragraph()
    set_rtl_paragraph(p)
    p.alignment = align
    pf = p.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = 1.3
    pf.keep_with_next = keep
    if indent:
        pf.right_indent = Inches(indent)
    r = p.add_run(text)
    set_run(r, size=size, bold=bold)
    return p


def heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    set_rtl_paragraph(p)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    set_run(r, size=22 if level == 1 else 18, bold=True)
    return p


def bullet(doc, text):
    return paragraph(doc, f"•  {text}", size=16, after=6, indent=0.18)


def numbered(doc, number, text):
    return paragraph(doc, f"{number}  {text}", size=16, after=7, indent=0.18)


def page_break(doc):
    doc.add_page_break()


def set_cell_shading(cell, fill):
    tcpr = cell._tc.get_or_add_tcPr()
    shd = tcpr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tcpr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=120, start=140, bottom=120, end=140):
    tcpr = cell._tc.get_or_add_tcPr()
    mar = tcpr.first_child_found_in("w:tcMar")
    if mar is None:
        mar = OxmlElement("w:tcMar")
        tcpr.append(mar)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table):
    tblpr = table._tbl.tblPr
    borders = tblpr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tblpr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        elem = OxmlElement(f"w:{edge}")
        elem.set(qn("w:val"), "single")
        elem.set(qn("w:sz"), "8")
        elem.set(qn("w:color"), LIGHT_GRAY)
        borders.append(elem)


def make_rtl_table(table, widths):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    tblpr = table._tbl.tblPr
    bidi_visual = OxmlElement("w:bidiVisual")
    bidi_visual.set(qn("w:val"), "1")
    tblpr.append(bidi_visual)
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tblpr.append(layout)
    total_twips = int(sum(widths) * 1440)
    width = OxmlElement("w:tblW")
    width.set(qn("w:w"), str(total_twips))
    width.set(qn("w:type"), "dxa")
    tblpr.append(width)
    for col, w in zip(table.columns, widths):
        col.width = Inches(w)
    set_table_borders(table)


def fill_table_cell(cell, text, *, size=14.5, bold=False, color=BLACK, center=False):
    set_cell_margins(cell)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    p = cell.paragraphs[0]
    set_rtl_paragraph(p)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.2
    r = p.add_run(text)
    set_run(r, size=size, bold=bold, color=color)


def add_page_number(paragraph):
    set_rtl_paragraph(paragraph)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    visible = OxmlElement("w:t")
    visible.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for element in (begin, instr, separate, visible, end):
        run._r.append(element)
    set_run(run, size=11, color="666666")


def toc_line(doc, label, page):
    p = paragraph(doc, "", size=15, after=5)
    p.paragraph_format.tab_stops.add_tab_stop(Inches(0.35))
    r = p.add_run(f"{label}  ....................................................................  {page}")
    set_run(r, size=15)


doc = Document()
section = doc.sections[0]
section.page_width = Cm(21)
section.page_height = Cm(29.7)
section.top_margin = Cm(2)
section.bottom_margin = Cm(1.8)
section.left_margin = Cm(2)
section.right_margin = Cm(2)
section.different_first_page_header_footer = True

sect_bidi = OxmlElement("w:bidi")
sect_bidi.set(qn("w:val"), "1")
section._sectPr.append(sect_bidi)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = FONT
normal.font.size = Pt(16)
normal.font.color.rgb = RGBColor(0, 0, 0)
normal.paragraph_format.space_after = Pt(8)
normal.paragraph_format.line_spacing = 1.3
normal_ppr = normal.element.get_or_add_pPr()
normal_bidi = OxmlElement("w:bidi")
normal_bidi.set(qn("w:val"), "1")
normal_ppr.append(normal_bidi)

title_style = styles["Title"]
title_style.font.name = FONT
title_style.font.size = Pt(28)
title_style.font.bold = True
title_style.font.color.rgb = RGBColor(0, 0, 0)
remove_paragraph_borders(title_style.element.get_or_add_pPr())

for name, size, before, after in (("Heading 1", 22, 12, 7), ("Heading 2", 18, 8, 4)):
    style = styles[name]
    style.font.name = FONT
    style.font.size = Pt(size)
    style.font.bold = True
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.paragraph_format.space_before = Pt(before)
    style.paragraph_format.space_after = Pt(after)
    style.paragraph_format.keep_with_next = True
    ppr = style.element.get_or_add_pPr()
    remove_paragraph_borders(ppr)
    style_bidi = OxmlElement("w:bidi")
    style_bidi.set(qn("w:val"), "1")
    ppr.append(style_bidi)

add_page_number(section.footer.paragraphs[0])

# Page 1: cover
for line, size in (
    ("المملكة العربية السعودية", 15),
    ("وزارة التعليم", 15),
    ("إدارة التعليم بالمنطقة الشرقية", 14),
    ("مكتب تعليم القطيف", 14),
    ("مدرسة دار الحكمة الثانوية", 14),
):
    paragraph(doc, line, size=size, bold=True, align=WD_ALIGN_PARAGRAPH.LEFT, after=2)

for _ in range(3):
    doc.add_paragraph()

title = doc.add_paragraph(style="Title")
set_rtl_paragraph(title)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
title.paragraph_format.space_after = Pt(14)
remove_paragraph_borders(title._p.get_or_add_pPr())
set_run(title.add_run("جمع البيانات وتحليلها لتصميم تطبيق دراستي لتنظيم الواجبات والمذاكرة"), size=28, bold=True)

paragraph(doc, "بحث يوضح كيف تساعد بيانات الطلاب في تحديد خصائص التطبيق وتصميمه واختباره وتطويره", size=17, align=WD_ALIGN_PARAGRAPH.CENTER, after=24)

for _ in range(2):
    doc.add_paragraph()

paragraph(doc, "أُعد هذا البحث كمتطلب لمادة المهارات الرقمية والحاسب الآلي", size=16, align=WD_ALIGN_PARAGRAPH.CENTER, after=18)
paragraph(doc, "عمل الطالب  ............................................................", size=16, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, after=9)
paragraph(doc, "الشعبة  ٣٠١", size=16, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, after=9)
paragraph(doc, "الرقم في كشف الدرجات  ........................", size=16, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)

page_break(doc)

# Page 2: dedication
heading(doc, "إهداء", 1)
paragraph(doc, "إلى أمي وأبي،", size=19, bold=True, before=20, after=12)
paragraph(doc, "أهديكما هذا العمل تقديرًا لدعمكما وتشجيعكما المستمرين. فقد كان حرصكما على تعليمي وتنظيم وقتي دافعًا لي لاختيار فكرة تطبيق يساعد الطلاب على متابعة واجباتهم ومذاكرتهم بصورة أوضح وأسهل.", size=17, after=10)

page_break(doc)

# Page 3: introduction and app definition
heading(doc, "المقدمة", 1)
paragraph(doc, "يواجه كثير من الطلاب صعوبة في متابعة الواجبات ومواعيد الاختبارات وتوزيع وقت المذاكرة بين المواد. وقد يؤدي الاعتماد على الذاكرة أو الأوراق المتفرقة إلى نسيان بعض المهام أو تنفيذها في وقت متأخر. لذلك يقترح هذا البحث تصميم تطبيق محدد باسم دراستي يساعد الطالب على تسجيل واجباته، وترتيب أولوياته، ووضع خطة مذاكرة يومية، واستقبال تنبيهات قبل موعد التسليم أو الاختبار.")
paragraph(doc, "لا يهدف البحث إلى كتابة التطبيق مباشرة، بل إلى توضيح كيف يمكن جمع بيانات الطلاب وتحليلها قبل التطوير وأثناءه. وتساعد هذه البيانات على معرفة المشكلات الحقيقية التي يواجهها الطالب، وتحديد الخصائص الضرورية، ثم اختبار سهولة الاستخدام وتحسين التطبيق بناءً على نتائج واضحة بدلًا من التخمين.")

heading(doc, "تعريف تطبيق دراستي", 1)
paragraph(doc, "دراستي تطبيق تعليمي مقترح للهواتف الذكية، مخصص لطلاب المرحلة الثانوية. يجمع في مكان واحد جدول المواد، وقائمة الواجبات، ومواعيد الاختبارات، وخطة المذاكرة، والتنبيهات اليومية. ويعتمد نجاحه على بساطة الواجهة وسرعة إدخال المهمة ووضوح ما يجب إنجازه في كل يوم.")

page_break(doc)

# Page 4: problem, goals, users
heading(doc, "مشكلة البحث", 1)
paragraph(doc, "تتمثل المشكلة في أن الطالب قد يستخدم أكثر من وسيلة لمتابعة دراسته، مثل الرسائل والمنصات التعليمية والدفاتر والتقويم. هذا التشتت يجعل الوصول إلى الواجب المطلوب أو موعد الاختبار أبطأ، وقد يزيد احتمال النسيان. ومن هنا تظهر الحاجة إلى تطبيق واحد ينظم المعلومات الدراسية ويعرضها حسب الأولوية والوقت المتبقي.")

heading(doc, "أهداف تطبيق دراستي", 1)
bullet(doc, "تجميع الواجبات والاختبارات وخطط المذاكرة في واجهة واحدة.")
bullet(doc, "تنبيه الطالب قبل المواعيد المهمة بوقت يحدده بنفسه.")
bullet(doc, "ترتيب المهام حسب تاريخ التسليم ودرجة الأهمية.")
bullet(doc, "إظهار التقدم اليومي والأسبوعي بطريقة سهلة الفهم.")
bullet(doc, "تقليل البيانات الشخصية التي يطلبها التطبيق وحماية ما يُجمع منها.")

heading(doc, "المستخدمون المستهدفون", 1)
paragraph(doc, "المستخدم الأساسي هو طالب المرحلة الثانوية، ويمكن أن يستفيد من التطبيق أيضًا ولي الأمر أو المعلم إذا أضيفت لاحقًا خاصية مشاركة اختيارية. ويجب أن يبقى الطالب قادرًا على استخدام الخصائص الأساسية من دون مشاركة بياناته مع أي طرف آخر.")

page_break(doc)

# Page 5: data collection
heading(doc, "طرق جمع البيانات لتطبيق دراستي", 1)
heading(doc, "الاستبيان الرقمي", 2)
paragraph(doc, "يوزع استبيان قصير على عينة من طلاب المرحلة الثانوية. ويسأل عن الوسائل التي يستخدمونها حاليًا، وعدد المرات التي ينسون فيها مهمة، ونوع التنبيه المفضل، وأكثر الخصائص التي يحتاجون إليها. وتُستخدم أسئلة اختيار من متعدد مع سؤال مفتوح يسمح للطالب بوصف مشكلة لم تتناولها الخيارات.")

heading(doc, "المقابلات الفردية", 2)
paragraph(doc, "تُجرى مقابلات قصيرة مع عدد محدود من الطلاب لفهم تفاصيل يومهم الدراسي، وكيف يسجلون الواجبات، ومتى يراجعون جدولهم. وتفيد المقابلة في اكتشاف أسباب السلوك، مثل تجاهل التنبيهات الكثيرة أو صعوبة إدخال الواجب بسرعة.")

heading(doc, "ملاحظة تجربة الاستخدام", 2)
paragraph(doc, "يُطلب من الطالب تنفيذ مهام على نموذج أولي، مثل إضافة واجب وتعديل موعد اختبار ووضع علامة إتمام. ويُسجل الوقت المطلوب لكل مهمة، والخطوات التي سببت ترددًا، والأخطاء التي حدثت، من دون تصوير الطالب أو تسجيل بيانات لا يحتاج إليها الاختبار.")

heading(doc, "بيانات الاستخدام بعد الإطلاق", 2)
paragraph(doc, "بعد إطلاق نسخة تجريبية، يمكن قياس عدد المهام التي أُضيفت وأُكملت، والخصائص الأكثر استخدامًا، ونسبة فتح التنبيهات. ويجب أن تكون هذه القياسات مجمعة قدر الإمكان، مع توضيح سياسة الخصوصية وإتاحة إيقاف التحليلات الاختيارية.")

page_break(doc)

# Page 6: analysis and translating results
heading(doc, "طريقة تحليل البيانات", 1)
paragraph(doc, "تُراجع إجابات الاستبيان أولًا للتأكد من اكتمالها، ثم تُقسم إلى مجموعات قابلة للمقارنة. ويمكن حساب تكرار كل مشكلة ونسبة اختيار كل خاصية، بينما تُصنف إجابات المقابلات إلى موضوعات مثل النسيان، وكثرة المنصات، وصعوبة ترتيب الأولويات، والإزعاج من التنبيهات.")
paragraph(doc, "بعد ذلك تُقارن النتائج من أكثر من مصدر. فإذا ذكر الطلاب في الاستبيان حاجتهم إلى إدخال سريع، وظهر في الملاحظة أن إضافة الواجب تستغرق خطوات كثيرة، تصبح خاصية الإدخال المختصر أولوية واضحة. أما الآراء الفردية التي لا تتكرر فتُحفظ للمراجعة ولا تُعامل وحدها على أنها نتيجة عامة.")

heading(doc, "تحويل النتائج إلى قرارات تصميمية", 1)
table = doc.add_table(rows=1, cols=3)
make_rtl_table(table, [1.65, 2.25, 2.55])
headers = ("البيان الملحوظ", "ما يعنيه", "القرار في التطبيق")
for cell, text in zip(table.rows[0].cells, headers):
    set_cell_shading(cell, NAVY)
    fill_table_cell(cell, text, size=14, bold=True, color="FFFFFF", center=True)

rows = (
    ("نسيان مواعيد التسليم", "الحاجة إلى تذكير واضح", "تنبيه مرن قبل الموعد"),
    ("تعدد مصادر الواجبات", "صعوبة جمع المهام", "شاشة موحدة للواجبات"),
    ("كثرة الخطوات", "بطء تسجيل المهمة", "زر إضافة سريع"),
    ("تجاهل التنبيهات", "التنبيه غير مناسب", "اختيار الوقت ونوع التنبيه"),
)
for idx, row in enumerate(rows):
    cells = table.add_row().cells
    if idx % 2:
        for cell in cells:
            set_cell_shading(cell, PALE_BLUE)
    for cell, text in zip(cells, row):
        fill_table_cell(cell, text, size=13.5, center=False)

page_break(doc)

# Page 7: requirements
heading(doc, "متطلبات تطبيق دراستي", 1)
heading(doc, "المتطلبات الوظيفية", 2)
bullet(doc, "إضافة واجب أو اختبار مع المادة والتاريخ والوقت ودرجة الأولوية.")
bullet(doc, "عرض مهام اليوم والأسبوع وترتيبها حسب الموعد.")
bullet(doc, "إرسال تنبيه واحد أو أكثر وفق اختيار الطالب.")
bullet(doc, "وضع علامة على المهمة المكتملة وإظهار نسبة الإنجاز.")
bullet(doc, "إنشاء خطة مذاكرة قابلة للتعديل لكل مادة.")

heading(doc, "متطلبات سهولة الاستخدام", 2)
bullet(doc, "إتمام إضافة الواجب في خطوات قليلة وبأزرار واضحة.")
bullet(doc, "استخدام خط عربي كبير وتباين واضح بين النص والخلفية.")
bullet(doc, "تجنب ازدحام الشاشة بالمعلومات غير الضرورية.")
bullet(doc, "إتاحة تعديل المهمة أو حذفها من مكان ظاهر.")

heading(doc, "متطلبات الخصوصية والأمان", 2)
bullet(doc, "جمع أقل قدر ممكن من البيانات الشخصية.")
bullet(doc, "حماية حساب المستخدم وعدم إظهار مهامه للآخرين دون موافقته.")
bullet(doc, "شرح سبب جمع أي بيانات تحليلية وإتاحة التحكم فيها.")

page_break(doc)

# Page 8: workflow
heading(doc, "مراحل تصميم التطبيق وتطويره", 1)
numbered(doc, "١", "التخطيط  تُحدد مشكلة تنظيم الدراسة والفئة المستهدفة والأسئلة التي يجب أن تجيب عنها البيانات.")
numbered(doc, "٢", "جمع البيانات  تُستخدم الاستبيانات والمقابلات وملاحظة الاستخدام لفهم احتياجات الطلاب.")
numbered(doc, "٣", "تصميم النموذج الأولي  تُرسم الشاشات الأساسية، ومنها الصفحة الرئيسية وإضافة المهمة والتقويم وخطة المذاكرة.")
numbered(doc, "٤", "اختبار سهولة الاستخدام  ينفذ الطلاب مهام محددة، وتُسجل الصعوبات والوقت والأخطاء.")
numbered(doc, "٥", "التطوير البرمجي  تُبنى الخصائص التي ثبتت أهميتها، مع تطبيق ضوابط الخصوصية والأمان.")
numbered(doc, "٦", "الإطلاق التجريبي  يُتاح التطبيق لمجموعة صغيرة، وتُقاس الأعطال واستخدام الخصائص والاستفادة من التنبيهات.")
numbered(doc, "٧", "التحسين المستمر  تُراجع البيانات وآراء الطلاب بصورة دورية، ثم تُعدل الأولويات والواجهات.")

heading(doc, "اختبار نجاح التطبيق", 1)
paragraph(doc, "يُعد التطبيق ناجحًا إذا استطاع الطالب إضافة واجب والعثور عليه وتحديد موعده دون مساعدة، وإذا انخفض عدد المهام المنسية لدى المستخدمين التجريبيين، وإذا كانت التنبيهات مفيدة وغير مزعجة. كما يجب ألا تحدث أعطال متكررة أو تسربات للبيانات.")

page_break(doc)

# Page 9: expected results and recommendations
heading(doc, "النتائج المتوقعة", 1)
bullet(doc, "أن تكون شاشة المهام اليومية هي الشاشة الأكثر أهمية للطلاب.")
bullet(doc, "أن يفضل الطلاب الإدخال السريع على النماذج الطويلة.")
bullet(doc, "أن تختلف أوقات التنبيه المناسبة من طالب إلى آخر، ولذلك يجب أن تكون قابلة للتخصيص.")
bullet(doc, "أن يزيد التزام الطالب بالخطة عندما يرى المهام المكتملة والتقدم الأسبوعي.")
bullet(doc, "أن يرفع وضوح سياسة الخصوصية وثقة المستخدم في التطبيق.")

heading(doc, "الاستنتاج", 1)
paragraph(doc, "توضح دراسة تطبيق دراستي أن جمع البيانات ليس مرحلة منفصلة عن التصميم، بل هو أساس تحديد المشكلة والخصائص وترتيب الأولويات. كما أن الجمع بين الاستبيان والمقابلة واختبار الاستخدام يعطي صورة أفضل من الاعتماد على وسيلة واحدة. وتبقى النتائج المتوقعة بحاجة إلى اختبار فعلي على عينة من الطلاب قبل اعتمادها نهائيًا.")

heading(doc, "التوصيات", 1)
bullet(doc, "بدء التطوير بالخصائص الأساسية فقط، ثم إضافة الخصائص الجديدة بعد إثبات الحاجة إليها.")
bullet(doc, "إشراك طلاب من مستويات دراسية مختلفة في الاختبار.")
bullet(doc, "تجربة التطبيق على أحجام شاشات مختلفة والتأكد من وضوح الخط العربي.")
bullet(doc, "مراجعة التنبيهات دوريًا حتى تساعد الطالب ولا تتحول إلى مصدر إزعاج.")

page_break(doc)

# Page 10: conclusion and references
heading(doc, "الخاتمة", 1)
paragraph(doc, "يقدم تطبيق دراستي حلًا محددًا لمشكلة واقعية هي تشتت الواجبات ومواعيد الاختبارات وخطط المذاكرة. ويبين هذا البحث أن الطريق إلى تطبيق مفيد يبدأ بفهم الطالب، ثم تحويل البيانات إلى متطلبات واضحة، واختبار التصميم قبل التوسع في البرمجة. وعندما تُجمع البيانات بموافقة المستخدم وتُحمى بصورة مناسبة، يمكن الاستفادة منها لتحسين التطبيق من دون الإضرار بخصوصيته.")

heading(doc, "المراجع", 1)
bullet(doc, "الهيئة السعودية للبيانات والذكاء الاصطناعي  مواد التوعية بنظام حماية البيانات الشخصية.")
bullet(doc, "مركز مساعدة نماذج جوجل  إرشادات إنشاء الاستبيانات وتنظيم الردود.")
bullet(doc, "وثائق تحليلات فايربيس  قياس استخدام التطبيقات والأحداث داخلها.")
bullet(doc, "إرشادات ماتيريال ديزاين  مبادئ تصميم واجهات التطبيقات وسهولة الاستخدام.")

page_break(doc)

# Page 11: contents on its own page
heading(doc, "الفهرس", 1)
paragraph(doc, "", size=12, after=5)
for label, page in (
    ("إهداء", "٢"),
    ("المقدمة", "٣"),
    ("تعريف تطبيق دراستي", "٣"),
    ("مشكلة البحث", "٤"),
    ("أهداف تطبيق دراستي", "٤"),
    ("المستخدمون المستهدفون", "٤"),
    ("طرق جمع البيانات لتطبيق دراستي", "٥"),
    ("طريقة تحليل البيانات", "٦"),
    ("تحويل النتائج إلى قرارات تصميمية", "٦"),
    ("متطلبات تطبيق دراستي", "٧"),
    ("مراحل تصميم التطبيق وتطويره", "٨"),
    ("اختبار نجاح التطبيق", "٨"),
    ("النتائج المتوقعة", "٩"),
    ("الاستنتاج", "٩"),
    ("التوصيات", "٩"),
    ("الخاتمة", "١٠"),
    ("المراجع", "١٠"),
):
    toc_line(doc, label, page)

props = doc.core_properties
props.title = "جمع البيانات وتحليلها لتصميم تطبيق دراستي لتنظيم الواجبات والمذاكرة"
props.subject = "بحث مدرسي عن تصميم تطبيق دراستي بالاعتماد على البيانات"
props.author = ""
props.last_modified_by = ""

doc.save(OUT)
print(OUT)
