from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT = Path(r"C:\Users\WIN-10\Documents\Abqoor\Research_Data_Analysis_Applications.docx")
FONT_AR = "Arial"
BLACK = "000000"
NAVY = "17365D"
PALE_BLUE = "EAF2F8"
LIGHT_GRAY = "D9D9D9"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color=LIGHT_GRAY, size="8"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:color"), color)


def set_run_font(run, size=None, bold=None, color=BLACK, rtl=True):
    run.font.name = FONT_AR
    run.font.color.rgb = RGBColor.from_string(color)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.get_or_add_rFonts()
    for attr in ("ascii", "hAnsi", "cs", "eastAsia"):
        r_fonts.set(qn(f"w:{attr}"), FONT_AR)
    if rtl:
        rtl_el = r_pr.find(qn("w:rtl"))
        if rtl_el is None:
            rtl_el = OxmlElement("w:rtl")
            r_pr.append(rtl_el)
        rtl_el.set(qn("w:val"), "1")
    lang = r_pr.find(qn("w:lang"))
    if lang is None:
        lang = OxmlElement("w:lang")
        r_pr.append(lang)
    lang.set(qn("w:bidi"), "ar-SA")


def set_rtl(paragraph, enabled=True):
    p_pr = paragraph._p.get_or_add_pPr()
    bidi = p_pr.find(qn("w:bidi"))
    if enabled and bidi is None:
        bidi = OxmlElement("w:bidi")
        p_pr.append(bidi)
    if bidi is not None:
        bidi.set(qn("w:val"), "1" if enabled else "0")


def add_text(doc, text, *, size=14, bold=False, align=WD_ALIGN_PARAGRAPH.RIGHT,
             before=0, after=7, rtl=True, keep_with_next=False):
    p = doc.add_paragraph()
    p.alignment = align
    set_rtl(p, rtl)
    fmt = p.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = 1.25
    fmt.keep_with_next = keep_with_next
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold, rtl=rtl)
    return p


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_rtl(p)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    set_run_font(run, size=18 if level == 1 else 15, bold=True)
    return p


def add_bullet(doc, text):
    p = add_text(doc, f"•  {text}", size=13.5, after=5)
    p.paragraph_format.right_indent = Inches(0.18)
    p.paragraph_format.left_indent = Inches(0.05)
    return p


def add_numbered(doc, number, text):
    p = add_text(doc, f"{number}.  {text}", size=13.5, after=6)
    p.paragraph_format.right_indent = Inches(0.18)
    return p


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for node in (begin, instr, separate, text, end):
        run._r.append(node)
    set_run_font(run, size=10, rtl=False, color="666666")


def add_toc_line(doc, label, page):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_rtl(p)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.15
    tabs = p.paragraph_format.tab_stops
    tabs.add_tab_stop(Inches(0.35))
    run = p.add_run(f"{label}  ........................................................................  {page}")
    set_run_font(run, size=12.5)


doc = Document()
section = doc.sections[0]
section.top_margin = Inches(0.8)
section.bottom_margin = Inches(0.75)
section.left_margin = Inches(0.9)
section.right_margin = Inches(0.9)
section.different_first_page_header_footer = True

sect_pr = section._sectPr
bidi = OxmlElement("w:bidi")
bidi.set(qn("w:val"), "1")
sect_pr.append(bidi)

normal = doc.styles["Normal"]
normal.font.name = FONT_AR
normal.font.size = Pt(13.5)
normal.font.color.rgb = RGBColor(0, 0, 0)
normal.paragraph_format.space_after = Pt(7)
normal.paragraph_format.line_spacing = 1.25

title_style = doc.styles["Title"]
title_style.font.name = FONT_AR
title_style.font.size = Pt(24)
title_style.font.bold = True
title_style.font.color.rgb = RGBColor(0, 0, 0)
title_ppr = title_style.element.get_or_add_pPr()
title_border = title_ppr.find(qn("w:pBdr"))
if title_border is not None:
    title_ppr.remove(title_border)

for name, size, before, after in (("Heading 1", 18, 14, 7), ("Heading 2", 15, 9, 4)):
    style = doc.styles[name]
    style.font.name = FONT_AR
    style.font.size = Pt(size)
    style.font.bold = True
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.paragraph_format.space_before = Pt(before)
    style.paragraph_format.space_after = Pt(after)
    style.paragraph_format.keep_with_next = True

footer = section.footer
footer.is_linked_to_previous = False
add_page_number(footer.paragraphs[0])

# Cover page
for line, size in (
    ("المملكة العربية السعودية", 13),
    ("وزارة التعليم", 13),
    ("إدارة التعليم بالمنطقة الشرقية", 12),
    ("مكتب تعليم القطيف", 12),
    ("مدرسة دار الحكمة الثانوية", 12),
):
    add_text(doc, line, size=size, bold=True, align=WD_ALIGN_PARAGRAPH.LEFT, after=2)

for _ in range(4):
    doc.add_paragraph()

title = doc.add_paragraph(style="Title")
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
set_rtl(title)
title.paragraph_format.space_after = Pt(10)
direct_border = title._p.get_or_add_pPr().find(qn("w:pBdr"))
if direct_border is not None:
    title._p.get_or_add_pPr().remove(direct_border)
run = title.add_run("طرق تحليل وجمع البيانات واستخدامها في عمل تطبيق خاص")
set_run_font(run, size=24, bold=True)

subtitle = add_text(
    doc,
    "بحث في دور البيانات خلال مراحل تخطيط التطبيقات وتصميمها واختبارها وتحسينها",
    size=14,
    align=WD_ALIGN_PARAGRAPH.CENTER,
    after=24,
)

for _ in range(2):
    doc.add_paragraph()

add_text(doc, "أُعد هذا البحث كمتطلب لمادة المهارات الرقمية والحاسب الآلي", size=13.5, align=WD_ALIGN_PARAGRAPH.CENTER, after=16)
add_text(doc, "عمل الطالب  ............................................................", size=13.5, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, after=8)
add_text(doc, "الشعبة  ٣٠١", size=13.5, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, after=8)
add_text(doc, "الرقم في كشف الدرجات  ........................", size=13.5, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)

doc.add_page_break()

# Dedication
add_heading(doc, "إهداء", 1)
add_text(doc, "إلى أمي وأبي،", size=16, bold=True, before=18, after=10)
add_text(doc, "لكما كل الامتنان والتقدير على دعمي وتشجيعي المستمرين. لقد كنتما دائمًا مصدر إلهامي ودافعي لتحقيق النجاح. هذا العمل ثمرة جهودكما ومحبتكما التي لا تنضب.", size=14, after=8)

doc.add_page_break()

# Introduction and rationale
add_heading(doc, "المقدمة", 1)
add_text(doc, "أصبحت التطبيقات الإلكترونية من أهم الوسائل المستخدمة في حياتنا اليومية؛ فهي تساعد في التعليم والتواصل والتسوق وإدارة الأعمال والترفيه وتنظيم الوقت. ومع زيادة الاعتماد على التطبيقات، أصبح من الضروري تصميم تطبيقات تلبي احتياجات المستخدمين وتوفر لهم تجربة سهلة ومفيدة. ولكي يُنشأ تطبيق ناجح، لا تكفي جودة الفكرة وحدها، بل يجب معرفة المستخدمين المستهدفين واحتياجاتهم والمشكلات التي يواجهونها.")
add_text(doc, "وتُعد عملية جمع البيانات وتحليلها من أهم المراحل التي تساعد على نجاح التطبيق. فمن خلال البيانات يمكن معرفة ما يفضله المستخدمون، والمشكلات التي تواجههم، والميزات التي يحتاجون إليها. وبعد جمع البيانات تُحلل النتائج وتُستخدم في بناء التطبيق وهندسته ليعمل بكفاءة ويوفر تجربة استخدام مناسبة.")

add_heading(doc, "دور جمع وتحليل البيانات في تطوير التطبيقات", 1)
add_text(doc, "يسهم فهم متطلبات المستخدمين من خلال البيانات في تحسين القرارات التصميمية والبرمجية وتقليل مخاطر بناء ميزات لا يحتاج إليها الجمهور المستهدف.")
add_bullet(doc, "تؤثر جودة البيانات التي تُجمع في مرحلة التخطيط في دقة تصميم واجهة الاستخدام وتجربة المستخدم.")
add_bullet(doc, "يساعد تحليل سلوك المستخدمين قبل التطوير وبعده على رفع فرص استمرارهم في استخدام التطبيق.")
add_bullet(doc, "تساعد البيانات في تحديد الحد الأدنى من المنتج القابل للاستخدام MVP، مما يقلل هدر الوقت في ميزات غير أساسية.")

doc.add_page_break()

# Methods and workflow
add_heading(doc, "طرق وأدوات جمع البيانات للتطبيقات", 1)
add_heading(doc, "١  الاستبيانات الرقمية", 2)
add_text(doc, "تُعد أداة سريعة لجمع بيانات كمية ونوعية من شريحة واسعة من المستخدمين المستهدفين قبل بدء المبادرة البرمجية. وتكون أكثر فائدة عندما تُصاغ أسئلتها بوضوح وتُربط بهدف محدد.")
add_heading(doc, "٢  المقابلات الشخصية ومجموعات التركيز", 2)
add_text(doc, "توفر فهمًا عميقًا وسياقيًا للمشكلات وسلوك الأفراد وتوقعاتهم من التطبيق الجديد، كما تتيح طرح أسئلة متابعة تكشف أسباب الإجابات.")
add_heading(doc, "٣  أدوات التحليل الرقمية", 2)
add_text(doc, "تتبع سلوك المستخدم على النماذج الأولية أو الإصدارات التجريبية، مثل العناصر الأكثر نقرًا ومسارات التنقل والأقسام الأكثر استخدامًا، مع مراعاة الخصوصية وأخذ الموافقات اللازمة.")

add_heading(doc, "مراحل استخدام البيانات في بناء التطبيق الخاص", 1)
stages = [
    ("مرحلة التخطيط والتحليل", "تُستخدم البيانات المجمعة لتحديد المشكلة الأساسية التي سيعالجها التطبيق وتحديد الشريحة المستهدفة بدقة."),
    ("مرحلة التصميم والتطوير", "تُترجم النتائج إلى مخططات هيكلية Wireframes وواجهات تلبي تطلعات المستخدم وتسهل تنقله داخل التطبيق."),
    ("مرحلة الاختبار والتحسين", "تُجمع بيانات الأداء والأخطاء التقنية من مجموعة من المختبرين لإصلاح المشكلات قبل الإطلاق الرسمي وبعده."),
]
table = doc.add_table(rows=1, cols=2)
table.autofit = False
table.alignment = 1
tbl_pr = table._tbl.tblPr
layout = OxmlElement("w:tblLayout")
layout.set(qn("w:type"), "fixed")
tbl_pr.append(layout)
tbl_width = OxmlElement("w:tblW")
tbl_width.set(qn("w:w"), "9000")
tbl_width.set(qn("w:type"), "dxa")
tbl_pr.append(tbl_width)
table.columns[0].width = Inches(1.75)
table.columns[1].width = Inches(4.50)
set_table_borders(table)
hdr = table.rows[0].cells
hdr[0].width = Inches(1.75)
hdr[1].width = Inches(4.50)
for cell, text in zip(hdr, ("المرحلة", "استخدام البيانات")):
    set_cell_shading(cell, NAVY)
    set_cell_margins(cell)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_rtl(p)
    r = p.add_run(text)
    set_run_font(r, size=12.5, bold=True, color="FFFFFF")
for idx, (stage, use) in enumerate(stages):
    cells = table.add_row().cells
    cells[0].width = Inches(1.75)
    cells[1].width = Inches(4.50)
    if idx % 2:
        for cell in cells:
            set_cell_shading(cell, PALE_BLUE)
    for cell in cells:
        set_cell_margins(cell, top=110, bottom=110)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    for cell, text, bold_text in ((cells[0], stage, True), (cells[1], use, False)):
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        set_rtl(p)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.15
        r = p.add_run(text)
        set_run_font(r, size=12, bold=bold_text)

doc.add_page_break()

# Analysis and recommendations
add_heading(doc, "تحليل النتائج", 1)
add_bullet(doc, "الاعتماد على بيانات حقيقية أثناء تصميم التطبيقات يقلل الحاجة إلى إعادة العمل على أجزاء كبيرة بعد الإطلاق.")
add_bullet(doc, "لا يقتصر نجاح التطبيق على الفكرة البرمجية؛ بل يعتمد أيضًا على استجابة المطورين لاحتياجات المستخدمين وتحديث التطبيق بناءً على نتائج القياس.")
add_bullet(doc, "إهمال جمع البيانات وتحليلها قد يؤدي إلى تطبيقات معقدة أو صعبة الاستخدام أو بعيدة عن الغرض الأساسي منها.")

add_heading(doc, "الاستنتاج", 1)
add_bullet(doc, "جمع البيانات وتحليلها أساس مهم لبناء تطبيق إلكتروني ناجح وقابل للاستمرار.")
add_bullet(doc, "تمنح أدوات جمع البيانات وتحليلها المطورين رؤية أوضح لتحديد ميزات التطبيق وأولوياتها.")
add_bullet(doc, "يسهم التطوير القائم على البيانات في تحسين أداء التطبيق بصورة مستمرة ومواكبة تغير احتياجات المستخدمين.")

add_heading(doc, "التوصيات", 1)
add_numbered(doc, "١", "إجراء دراسة للسوق وجمع بيانات حقيقية قبل البدء في كتابة الشفرة البرمجية للتطبيق.")
add_numbered(doc, "٢", "اعتماد تصاميم بسيطة ومرنة للواجهات بناءً على تفضيلات المستخدمين التي يكشفها التحليل.")
add_numbered(doc, "٣", "إضافة وسيلة داخل التطبيق لجمع آراء المستخدمين بصورة منتظمة بعد الإطلاق لضمان التحسين المستمر.")
add_numbered(doc, "٤", "التركيز على أمن بيانات المستخدمين وخصوصيتها، وجمع الحد الضروري منها فقط لتعزيز الثقة في التطبيق.")

doc.add_page_break()

# Conclusion, references, and contents
add_heading(doc, "الخاتمة", 1)
add_text(doc, "يتجاوز إنتاج تطبيق ناجح امتلاك مهارات البرمجة وحدها؛ فهو يتطلب فهمًا عميقًا لاحتياجات المستخدمين وسلوكهم. وتساعد هندسة التطبيقات القائمة على البيانات على تقديم حلول رقمية أكثر فاعلية للمشكلات اليومية. لذلك ينبغي أن تكون آليات جمع البيانات وتحليلها جزءًا أصيلًا من دورة حياة تطوير البرمجيات، مع الالتزام بأمن البيانات وخصوصية أصحابها.")

add_heading(doc, "المراجع المقترحة", 1)
add_text(doc, "يمكن الاستفادة من قواعد البيانات العلمية الآتية للبحث عن دراسات محكّمة حول هندسة البرمجيات وتجربة المستخدم وتحليلات التطبيقات:", size=12.5)
for ref in (
    "IEEE Xplore Digital Library",
    "ACM Digital Library",
    "SpringerLink",
    "ScienceDirect",
):
    add_text(doc, f"•  {ref}", size=12.5, align=WD_ALIGN_PARAGRAPH.LEFT, after=4, rtl=False)

add_heading(doc, "الفهرس", 1)
for label, page in (
    ("إهداء", "٢"),
    ("المقدمة", "٣"),
    ("دور جمع وتحليل البيانات في تطوير التطبيقات", "٣"),
    ("طرق وأدوات جمع البيانات للتطبيقات", "٤"),
    ("مراحل استخدام البيانات في بناء التطبيق الخاص", "٤"),
    ("تحليل النتائج", "٥"),
    ("الاستنتاج", "٥"),
    ("التوصيات", "٥"),
    ("الخاتمة", "٦"),
    ("المراجع المقترحة", "٦"),
):
    add_toc_line(doc, label, page)

# Prevent Word from adding identifying metadata supplied by the local account.
props = doc.core_properties
props.title = "طرق تحليل وجمع البيانات واستخدامها في عمل تطبيق خاص"
props.subject = "بحث مدرسي عن استخدام البيانات في تطوير التطبيقات"
props.author = ""
props.last_modified_by = ""

doc.save(OUT)
print(OUT)
